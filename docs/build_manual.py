from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Manual_do_Usuario_SIGEA_SENAC.docx"
LOGO = ROOT / "Frontend" / "src" / "assets" / "images" / "logo-sigea.png"
TEACHERS_SHOT = Path(r"C:\Users\paulo\AppData\Local\Temp\codex-clipboard-f241ebf3-709e-4e57-ac33-195129e16851.png")
PASSWORD_SHOT = Path(r"C:\Users\paulo\AppData\Local\Temp\codex-clipboard-e561f3c5-ba6c-4f62-ad5a-7df2bb48857b.png")
TABLE_HELPER = Path(r"C:\Users\paulo\.codex\plugins\cache\openai-primary-runtime\documents\26.805.11740\skills\documents\scripts")
sys.path.insert(0, str(TABLE_HELPER))
from table_geometry import apply_table_geometry

NAVY = "0A4978"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
ORANGE = "F7941D"
INK = "172C3B"
MUTED = "62798A"
LIGHT_BLUE = "E8EEF5"
LIGHT_ORANGE = "FFF4E5"
LIGHT_GRAY = "F2F4F7"
GREEN = "16815A"
RED = "A33A3A"
WHITE = "FFFFFF"
CONTENT_DXA = 9360


def set_run(run, size=None, color=INK, bold=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color="D8E1E8", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_numbering(doc, fmt, text, left=540, hanging=270):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "singleLevel")
    abstract.append(multilevel)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def numbered_paragraph(doc, num_id, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), bold=True)
        set_run(p.add_run(text[len(bold_prefix):]))
    else:
        set_run(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_para(doc, text, bold_prefix=None, italic=False, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), bold=True)
        set_run(p.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        set_run(p.add_run(text), italic=italic)
    return p


def callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8" if edge == "left" else "3")
        border.set(qn("w:color"), accent if edge == "left" else "D8E1E8")
        border.set(qn("w:space"), "6")
        borders.append(border)
    p_pr.append(borders)
    set_run(p.add_run(f"{label}: "), bold=True, color=accent)
    set_run(p.add_run(text))
    return p


def style_table(table, header=True):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if header and r_idx == 0:
                shade(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    set_run(run, size=9.5, bold=(header and r_idx == 0), color=(DARK_BLUE if header and r_idx == 0 else INK))
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def add_picture(doc, path, width, caption, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(8)
    set_run(cap.add_run(caption), size=9, color=MUTED, italic=True)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run(run, size=9, color=MUTED)


def configure_body_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(hp.add_run("SIGEA SENAC  |  Manual do Usuário"), size=9, color=MUTED, bold=True)
    footer = section.footer
    ft = footer.paragraphs[0]
    ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(ft.add_run("Versão 1.0  •  Agosto de 2026  |  "), size=9, color=MUTED)
    add_page_number(ft)


def page_break(doc):
    doc.add_page_break()


doc = Document()
setup_styles(doc)
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

# Capa editorial
spacer = doc.add_paragraph()
spacer.paragraph_format.space_before = Pt(58)
logo_p = doc.add_paragraph()
logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo_shape = logo_p.add_run().add_picture(str(LOGO), width=Inches(4.5))
logo_shape._inline.docPr.set("descr", "Logotipo do Sistema Integrado de Gestão de Espaços Acadêmicos")
kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
kicker.paragraph_format.space_before = Pt(26)
kicker.paragraph_format.space_after = Pt(10)
set_run(kicker.add_run("GUIA OPERACIONAL"), size=11, color=ORANGE, bold=True)
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
set_run(title.add_run("Manual do Usuário"), size=30, color=NAVY, bold=True)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(30)
set_run(subtitle.add_run("Sistema Integrado de Gestão de Espaços Acadêmicos"), size=15, color=DARK_BLUE)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(54)
set_run(meta.add_run("SIGEA SENAC  •  Versão 1.0  •  Agosto de 2026"), size=11, color=MUTED, bold=True)
desc = doc.add_paragraph()
desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(desc.add_run("Orientações para consulta, cadastros, alocação de salas e segurança de acesso"), size=10, color=MUTED, italic=True)

body_section = doc.add_section(WD_SECTION.NEW_PAGE)
configure_body_section(body_section)
bullet_id = add_numbering(doc, "bullet", "•", left=540, hanging=270)
decimal_id = add_numbering(doc, "decimal", "%1.", left=540, hanging=270)

# 1
add_heading(doc, "1. Sobre o SIGEA", 1)
add_para(doc, "O SIGEA é o sistema de gestão dos espaços acadêmicos do SENAC. Ele reúne unidades, salas, turmas, professores, usuários e alocações em uma única interface, facilitando o planejamento diário e reduzindo conflitos de horário.")
callout(doc, "Objetivo central", "mostrar onde cada atividade acontecerá, em qual período e horário, mantendo uma visão confiável da disponibilidade das salas.", LIGHT_ORANGE, ORANGE)
add_heading(doc, "O que pode ser feito", 2)
for item in [
    "Consultar indicadores e ocupações na Visão geral.",
    "Cadastrar e localizar salas por unidade.",
    "Organizar turmas, datas, turnos, instrutores e coordenadores.",
    "Criar alocações recorrentes por dia da semana e impedir choques de horário.",
    "Cadastrar, editar e apagar professores.",
    "Administrar usuários, perfis e senhas provisórias.",
]: numbered_paragraph(doc, bullet_id, item)
add_heading(doc, "Perfis de acesso", 2)
roles = doc.add_table(rows=1, cols=3)
for idx, text in enumerate(["Perfil", "Pode realizar", "Restrições"]): roles.cell(0, idx).text = text
for row in [
    ("Administrador", "Gerencia todos os cadastros, usuários e senhas.", "Não pode apagar nem inativar a própria conta."),
    ("Coordenador", "Gerencia espaços, turmas, alocações e professores.", "Não administra usuários."),
    ("Consulta", "Visualiza as informações do sistema.", "Não inclui, altera ou exclui registros."),
]:
    cells = roles.add_row().cells
    for idx, value in enumerate(row): cells[idx].text = value
style_table(roles)
apply_table_geometry(roles, [1800, 3780, 3780], indent_dxa=120)

page_break(doc)

# 2
add_heading(doc, "2. Acesso e navegação", 1)
add_heading(doc, "Entrar no sistema", 2)
for item in [
    "Abra o endereço fornecido pela equipe responsável. No ambiente local, use http://localhost:5173.",
    "Informe o e-mail institucional e a senha recebida.",
    "Se a senha for provisória ou estiver vencida, crie uma nova senha antes de continuar.",
    "Após a autenticação, o sistema abre a Visão geral.",
]: numbered_paragraph(doc, decimal_id, item)
add_heading(doc, "Elementos da tela", 2)
for item in [
    "Menu lateral: alterna entre Visão geral, Salas, Turmas, Alocações, Professores e Usuários.",
    "Unidade: filtra os dados por polo ou unidade do SENAC.",
    "Data de referência: define o dia usado no mapa de ocupação.",
    "Pesquisa: localiza rapidamente registros na lista aberta.",
    "Rodapé do menu: mostra o usuário conectado, o atalho de senha e o botão de saída.",
]: numbered_paragraph(doc, bullet_id, item)
callout(doc, "Boa prática", "sempre selecione a unidade e a data corretas antes de analisar a disponibilidade ou criar uma alocação.")
add_heading(doc, "Sair com segurança", 2)
add_para(doc, "Use o ícone de saída no rodapé do menu. Em computadores compartilhados, não feche apenas a janela: encerre a sessão para impedir que outra pessoa utilize sua conta.")

page_break(doc)

# 3
add_heading(doc, "3. Visão geral", 1)
add_para(doc, "A Visão geral resume a situação das salas para a data selecionada. Os cartões mostram o total de salas, quantas estão ocupadas, quantas estão disponíveis e a taxa de ocupação.")
add_heading(doc, "Agenda do dia", 2)
add_para(doc, "A área Ocupações do dia lista as atividades previstas com horário, sala e unidade. O botão Ver todas abre o cadastro completo de alocações.")
add_heading(doc, "Pendências", 2)
add_para(doc, "O painel destaca turmas ativas que ainda não possuem uma alocação. Use Revisar turmas para localizar os casos e definir sala, datas, horário e dias da semana.")
callout(doc, "Leitura recomendada", "um número alto de pendências significa que existem turmas cadastradas sem sala ou sem agenda completa. A pendência deve ser analisada antes do início das aulas.", LIGHT_ORANGE, ORANGE)
add_heading(doc, "Filtros", 2)
for item in [
    "Todas as unidades: apresenta a situação consolidada.",
    "Unidade específica: mostra apenas os espaços e atividades daquela unidade.",
    "Data de referência: permite consultar o planejamento de qualquer dia.",
]: numbered_paragraph(doc, bullet_id, item)

page_break(doc)

# 4
add_heading(doc, "4. Salas e turmas", 1)
add_heading(doc, "Cadastro de salas", 2)
add_para(doc, "A tela Salas apresenta nome, unidade, tipo, capacidade, capacidade recomendada e status. A pesquisa aceita nome, código, tipo ou unidade.")
for item in [
    "Clique em Nova sala.",
    "Informe unidade, código, nome, tipo e capacidades.",
    "Confira os dados e clique em Salvar.",
]: numbered_paragraph(doc, decimal_id, item)
callout(doc, "Atenção", "a exclusão de uma sala pode remover também suas alocações vinculadas. Confirme o impacto antes de apagar.", LIGHT_ORANGE, ORANGE)
add_heading(doc, "Cadastro de turmas", 2)
add_para(doc, "A tela Turmas registra código, nome, período, turno, dias da semana, instrutor, coordenador, segmento e situação. Quando uma turma ainda não tem sala, o sistema mostra a ação Alocar agora.")
for item in [
    "Cadastre a turma com datas de início e término.",
    "Defina turno e dias da semana.",
    "Associe o professor quando o cadastro estiver disponível.",
    "Use Alocar agora ou abra Alocações para escolher a sala e o horário.",
]: numbered_paragraph(doc, decimal_id, item)
add_heading(doc, "Dados importados", 2)
add_para(doc, "A carga inicial foi organizada a partir das planilhas fornecidas. Registros importados devem ser revisados, especialmente nomes de instrutores, dias, horários e salas que estavam incompletos na origem.")

page_break(doc)

# 5
add_heading(doc, "5. Alocações e conflitos", 1)
add_para(doc, "Uma alocação liga uma atividade ou turma a uma sala, com vigência, intervalo de horário e dias da semana. Ela pode representar uma aula recorrente ou uma reserva avulsa.")
add_heading(doc, "Criar uma alocação", 2)
for item in [
    "Abra Alocações e clique em Nova alocação.",
    "Escolha a sala e, quando aplicável, a turma.",
    "Informe o título da atividade.",
    "Defina data inicial, data final, horário de início e horário de término.",
    "Marque os dias da semana em que a atividade ocorrerá.",
    "Salve e confirme que o registro aparece na lista.",
]: numbered_paragraph(doc, decimal_id, item)
add_heading(doc, "Detecção de conflito", 2)
callout(doc, "Proteção automática", "o SIGEA bloqueia uma nova alocação quando a mesma sala já está ocupada em datas, horários e dias da semana coincidentes.", LIGHT_BLUE, BLUE)
add_para(doc, "Quando houver conflito, leia a atividade indicada pelo sistema e ajuste pelo menos um dos elementos: sala, período, dias ou horário. Não contorne o aviso duplicando a turma.")
add_heading(doc, "Excluir uma alocação", 2)
add_para(doc, "Use o × ao lado do registro e confirme. A exclusão é definitiva e libera a sala para novas atividades naquele período.")

page_break(doc)

# 6
add_heading(doc, "6. Professores", 1)
add_para(doc, "O cadastro de professores reúne matrícula, nome, contato, área de atuação, especialidade, unidade, observações e status.")
add_picture(doc, TEACHERS_SHOT, 6.3, "Figura 1 - Lista de professores, pesquisa e ações de cadastro.", "Tela de professores do SIGEA com pesquisa, status e ações")
add_heading(doc, "Incluir ou editar", 2)
for item in [
    "Use Novo professor para criar um cadastro.",
    "Use o lápis para abrir os dados existentes e corrigi-los.",
    "Defina o status como Ativo ou Inativo e clique em Salvar.",
]: numbered_paragraph(doc, decimal_id, item)
add_heading(doc, "Professores importados", 2)
add_para(doc, "Matrículas iniciadas por IMP- foram geradas durante a importação das turmas. Esses registros precisam de conferência e podem ser editados com a matrícula real, e-mail, telefone, área e especialidade.")
callout(doc, "Exclusão", "o × apaga definitivamente o professor. As turmas existentes são preservadas, mas deixam de manter o vínculo direto com esse cadastro.", LIGHT_ORANGE, ORANGE)

page_break(doc)

# 7
add_heading(doc, "7. Usuários e permissões", 1)
add_para(doc, "A tela Usuários é exclusiva para administradores. Ela mostra nome, e-mail, matrícula, perfil, status e validade da senha.")
add_heading(doc, "Cadastrar um usuário", 2)
for item in [
    "Clique em Novo usuário.",
    "Informe nome, matrícula, e-mail e perfil.",
    "Defina uma senha provisória com pelo menos oito caracteres.",
    "Entregue a senha provisória ao usuário por um canal seguro.",
    "No primeiro acesso, o usuário será obrigado a criar a própria senha.",
]: numbered_paragraph(doc, decimal_id, item)
add_picture(doc, PASSWORD_SHOT, 4.7, "Figura 2 - Edição de usuário e definição de senha provisória.", "Janela de edição de usuário com perfil, status e campo de senha")
add_heading(doc, "Editar, redefinir ou apagar", 2)
for item in [
    "Lápis: altera nome, matrícula, e-mail, perfil, status ou atribui nova senha provisória.",
    "×: apaga definitivamente o usuário após confirmação.",
    "Proteção: o administrador conectado não pode apagar nem inativar a própria conta.",
]: numbered_paragraph(doc, bullet_id, item)

page_break(doc)

# 8
add_heading(doc, "8. Política de senhas", 1)
add_para(doc, "A política foi criada para reduzir o uso prolongado de credenciais e assegurar que apenas o titular conheça a senha definitiva.")
add_heading(doc, "Regras", 2)
for item in [
    "Todo novo cadastro recebe uma senha provisória.",
    "A troca é obrigatória no primeiro acesso.",
    "A senha definitiva é válida por 60 dias.",
    "Após o vencimento, o painel fica bloqueado até a renovação.",
    "Uma redefinição administrativa bloqueia imediatamente inclusive sessões já abertas.",
    "A nova senha deve ter maiúscula, minúscula, número, caractere especial e no mínimo oito caracteres.",
]: numbered_paragraph(doc, bullet_id, item)
add_heading(doc, "Trocar a própria senha", 2)
for item in [
    "Clique no ícone de chave no rodapé do menu.",
    "Informe a senha atual.",
    "Digite e confirme a nova senha.",
    "Clique em Salvar nova senha.",
]: numbered_paragraph(doc, decimal_id, item)
callout(doc, "Segurança", "não reutilize senhas pessoais, não envie senhas definitivas por mensagem e nunca compartilhe sua conta com outro colaborador.", LIGHT_ORANGE, ORANGE)

page_break(doc)

# 9
add_heading(doc, "9. Rotina recomendada", 1)
add_heading(doc, "Antes do início das atividades", 2)
for item in [
    "Revisar turmas ativas e seus períodos.",
    "Completar professores importados com matrícula IMP-.",
    "Resolver as pendências de alocação.",
    "Confirmar capacidade e tipo de cada sala.",
    "Verificar conflitos e alterações de calendário.",
]: numbered_paragraph(doc, bullet_id, item)
add_heading(doc, "Conferência diária", 2)
for item in [
    "Selecionar a unidade e a data do dia.",
    "Analisar a taxa de ocupação e as salas disponíveis.",
    "Conferir a Agenda do dia.",
    "Registrar reservas extraordinárias antes de liberar a sala.",
    "Encerrar a sessão ao finalizar o trabalho.",
]: numbered_paragraph(doc, decimal_id, item)
add_heading(doc, "Antes de excluir", 2)
callout(doc, "Checklist", "confirme o registro, verifique vínculos, valide se a exclusão é realmente necessária e leia a mensagem de confirmação. Exclusões definitivas não possuem botão de desfazer.", LIGHT_ORANGE, ORANGE)

page_break(doc)

# 10
add_heading(doc, "10. Solução de problemas", 1)
trouble = doc.add_table(rows=1, cols=2)
trouble.cell(0, 0).text = "Situação"
trouble.cell(0, 1).text = "Como resolver"
for situation, solution in [
    ("E-mail ou senha incorretos", "Confira o e-mail, a senha e o estado da tecla Caps Lock. Solicite redefinição ao administrador se necessário."),
    ("Redefinição obrigatória", "Use a senha provisória ou atual e crie uma nova senha que atenda a todos os requisitos."),
    ("Acesso somente de leitura", "O perfil Consulta não pode alterar registros. Solicite revisão do perfil ao administrador."),
    ("Conflito de horário", "Escolha outra sala ou ajuste período, dias ou horário da alocação."),
    ("Registro duplicado", "Pesquise o código ou a matrícula antes de cadastrar. Edite o registro existente quando for o mesmo item."),
    ("Dados não aparecem", "Limpe a pesquisa, selecione Todas as unidades e confira a data de referência."),
]:
    cells = trouble.add_row().cells
    cells[0].text, cells[1].text = situation, solution
style_table(trouble)
apply_table_geometry(trouble, [2700, 6660], indent_dxa=120)
add_heading(doc, "Quando acionar o suporte", 2)
add_para(doc, "Informe qual tela estava aberta, unidade, data, ação realizada e mensagem exibida. Não envie senhas. Se possível, anexe uma captura sem dados sensíveis.")

page_break(doc)

# 11
add_heading(doc, "11. Orientações para administração técnica", 1)
add_para(doc, "Esta seção é destinada à equipe responsável por manter o ambiente do SIGEA.")
for item in [
    "Banco de dados: MySQL, banco sigea_senac.",
    "Backend: API Node.js, Express, TypeScript e TypeORM.",
    "Frontend: React com Vite.",
    "Produção: usar migrações do TypeORM e manter DB_SYNC=false.",
    "Backup: realizar cópias regulares do banco antes de atualizações ou cargas em massa.",
    "Credenciais: manter senhas e segredos somente nos arquivos de ambiente protegidos; nunca inseri-los em manuais ou repositórios.",
]: numbered_paragraph(doc, bullet_id, item)
add_heading(doc, "Checklist de publicação", 2)
for item in [
    "Configurar o endereço do frontend permitido pela API.",
    "Definir um segredo forte para os tokens de autenticação.",
    "Alterar credenciais temporárias de instalação.",
    "Validar backup e restauração do MySQL.",
    "Testar login, perfis, troca de senha e bloqueio por expiração.",
    "Executar as validações npm run check no backend e npm run build no frontend.",
]: numbered_paragraph(doc, decimal_id, item)
callout(doc, "Encerramento", "o SIGEA depende da qualidade dos cadastros. Revisões periódicas de turmas, professores, salas e alocações mantêm o mapa de ocupação confiável.")

doc.core_properties.title = "Manual do Usuário - SIGEA SENAC"
doc.core_properties.subject = "Guia operacional do Sistema Integrado de Gestão de Espaços Acadêmicos"
doc.core_properties.author = "SIGEA SENAC"
doc.core_properties.keywords = "SIGEA, SENAC, salas, turmas, alocações, manual"
doc.save(OUT)
print(OUT)
