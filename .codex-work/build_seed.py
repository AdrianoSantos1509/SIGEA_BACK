import json
import math
import re
import sys
from datetime import date, datetime, time

from openpyxl import load_workbook


def text(value):
    return "" if value is None else str(value).strip()


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return None


def parse_time_range(value):
    match = re.search(r"(\d{1,2}:\d{2}).*?(\d{1,2}:\d{2})", text(value))
    return (match.group(1), match.group(2)) if match else (None, None)


def number(value, default=0):
    try:
        result = int(float(value))
        return result if math.isfinite(result) else default
    except (TypeError, ValueError):
        return default


buildings = [
    {"code": "ENNIUS", "name": "CEP Ennius Muniz", "location": "Distrito Federal"},
    {"code": "FAC", "name": "Faculdade Senac / Polo 713 Sul", "location": "Brasília - DF"},
    {"code": "SAO_SEBASTIAO", "name": "Polo São Sebastião", "location": "São Sebastião - DF"},
    {"code": "TAGUATINGA", "name": "CEP Taguatinga", "location": "Taguatinga - DF"},
]
rooms = []
classes = []


def add_rooms(ws, building_code, name_row, capacity_row, start_col=4):
    for col in range(start_col, ws.max_column + 1, 2):
        name = text(ws.cell(name_row, col).value).replace("\n", " ")
        if not name:
            continue
        capacity = number(ws.cell(capacity_row, col).value)
        room_type = "Laboratório" if re.search(r"lab|ti|beleza|moda|fotografia|podcast|redes|hardware|games", name, re.I) else "Sala de aula"
        if "audit" in name.lower():
            room_type = "Auditório"
        rooms.append({
            "buildingCode": building_code,
            "code": re.sub(r"\s+", " ", name),
            "name": re.sub(r"\s+", " ", name),
            "capacity": capacity,
            "recommendedCapacity": math.floor(capacity * 0.8) if capacity else 0,
            "type": room_type,
        })


map_path, ennius_path, tag_path, output_path = sys.argv[1:5]
map_book = load_workbook(map_path, read_only=False, data_only=False)
add_rooms(map_book["CEP ENNIUS MUNIZ"], "ENNIUS", 3, 4)
add_rooms(map_book["POLO FAC 713 SUL"], "FAC", 2, 3)
add_rooms(map_book["POLO - SÃO SEBASTIÃO"], "SAO_SEBASTIAO", 2, 3)

tag_book = load_workbook(tag_path, read_only=False, data_only=False)
add_rooms(tag_book["Mapa de Sala 2025 Tag"], "TAGUATINGA", 2, 3)

ennius_book = load_workbook(ennius_path, read_only=False, data_only=False)
ws = ennius_book["Dados da Turma"]
for row in range(2, ws.max_row + 1):
    code = text(ws.cell(row, 1).value)
    name = text(ws.cell(row, 2).value)
    if not code or not name:
        continue
    unit = text(ws.cell(row, 6).value)
    building_code = "FAC" if "38" in unit or "FAC" in unit.upper() else "ENNIUS"
    start_time, end_time = parse_time_range(ws.cell(row, 9).value)
    day_text = text(ws.cell(row, 8).value).lower()
    weekdays = []
    aliases = [("2ª", "SEG"), ("3ª", "TER"), ("4ª", "QUA"), ("5ª", "QUI"), ("6ª", "SEX"), ("sáb", "SAB")]
    if re.search(r"2ª\s+a\s+6ª", day_text):
        weekdays = ["SEG", "TER", "QUA", "QUI", "SEX"]
    elif re.search(r"2ª\s+a\s+5ª", day_text):
        weekdays = ["SEG", "TER", "QUA", "QUI"]
    elif re.search(r"2ª\s+a\s+sáb", day_text):
        weekdays = ["SEG", "TER", "QUA", "QUI", "SEX", "SAB"]
    else:
        weekdays = [label for token, label in aliases if token in day_text or label.lower() in day_text]
    classes.append({
        "buildingCode": building_code,
        "classroomCode": text(ws.cell(row, 14).value),
        "code": code,
        "name": name,
        "abbreviation": text(ws.cell(row, 3).value),
        "workload": number(ws.cell(row, 7).value),
        "segment": "",
        "type": "Turma",
        "startDate": iso_date(ws.cell(row, 4).value),
        "endDate": iso_date(ws.cell(row, 5).value),
        "shift": "Matutino" if start_time and start_time < "12:00" else "Vespertino" if start_time and start_time < "18:01" else "Noturno",
        "startTime": start_time,
        "endTime": end_time,
        "weekdays": weekdays,
        "students": 0,
        "status": text(ws.cell(row, 10).value) or "Em andamento",
        "instructor": text(ws.cell(row, 11).value),
        "coordinator": text(ws.cell(row, 12).value),
        "notes": text(ws.cell(row, 21).value),
    })

ws = tag_book["Em andamento"]
day_columns = [(8, "SEG"), (9, "TER"), (10, "QUA"), (11, "QUI"), (12, "SEX"), (13, "SAB")]
shift_times = {
    "matutino": ("08:00", "12:00"),
    "vespertino": ("14:00", "18:00"),
    "noturno": ("19:00", "22:00"),
}
for row in range(2, ws.max_row + 1):
    code = text(ws.cell(row, 2).value)
    name = text(ws.cell(row, 3).value)
    if not code or not name:
        continue
    unit_name = text(ws.cell(row, 1).value)
    normalized_unit = re.sub(r"[^A-Z0-9]+", "_", unit_name.upper()).strip("_") or "TAGUATINGA"
    building_code = "TAGUATINGA" if unit_name.upper() == "CEP" else normalized_unit
    if building_code not in {item["code"] for item in buildings}:
        buildings.append({"code": building_code, "name": unit_name, "location": "Distrito Federal"})
    shift = text(ws.cell(row, 15).value) or "Não informado"
    start_time, end_time = shift_times.get(shift.lower().strip(), (None, None))
    weekdays = [day for col, day in day_columns if text(ws.cell(row, col).value).upper() == "X"]
    classes.append({
        "buildingCode": building_code,
        "classroomCode": text(ws.cell(row, 16).value),
        "code": code,
        "name": name,
        "abbreviation": "",
        "workload": number(ws.cell(row, 4).value),
        "segment": text(ws.cell(row, 5).value),
        "type": text(ws.cell(row, 14).value) or "Turma",
        "startDate": iso_date(ws.cell(row, 6).value),
        "endDate": iso_date(ws.cell(row, 7).value),
        "shift": shift,
        "startTime": start_time,
        "endTime": end_time,
        "weekdays": weekdays,
        "students": 0,
        "status": text(ws.cell(row, 20).value) or "Em andamento",
        "instructor": text(ws.cell(row, 17).value),
        "coordinator": text(ws.cell(row, 18).value),
        "notes": text(ws.cell(row, 19).value),
    })

deduped_rooms = {}
for room in rooms:
    deduped_rooms[(room["buildingCode"], room["code"].lower())] = room

deduped_classes = {}
for item in classes:
    deduped_classes[item["code"]] = item

with open(output_path, "w", encoding="utf-8") as output:
    json.dump({"buildings": buildings, "rooms": list(deduped_rooms.values()), "classes": list(deduped_classes.values())}, output, ensure_ascii=False, indent=2)
