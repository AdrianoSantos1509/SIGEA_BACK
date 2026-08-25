import json
import os
import sys
from datetime import date, datetime

from openpyxl import load_workbook


def serial(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def meaningful_rows(ws, row_limit=160, col_limit=50):
    rows = []
    max_row = min(ws.max_row, row_limit)
    max_col = min(ws.max_column, col_limit)
    for row_idx in range(1, max_row + 1):
        values = [serial(ws.cell(row_idx, col_idx).value) for col_idx in range(1, max_col + 1)]
        if any(value not in (None, "") for value in values):
            while values and values[-1] in (None, ""):
                values.pop()
            rows.append({"row": row_idx, "values": values})
    return rows


def analyze(path):
    workbook = load_workbook(path, read_only=False, data_only=False)
    result = {"file": path, "sheets": []}
    for ws in workbook.worksheets:
        result["sheets"].append({
            "title": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "merged_ranges": [str(item) for item in list(ws.merged_cells.ranges)[:80]],
            "tables": list(ws.tables.keys()),
            "auto_filter": str(ws.auto_filter.ref) if ws.auto_filter.ref else None,
            "rows": meaningful_rows(ws),
        })
    return result


output_dir = os.path.join(".codex-work", "workbook-analysis")
os.makedirs(output_dir, exist_ok=True)
for input_path in sys.argv[1:]:
    report = analyze(input_path)
    name = os.path.splitext(os.path.basename(input_path))[0]
    with open(os.path.join(output_dir, f"{name}.json"), "w", encoding="utf-8") as output:
        json.dump(report, output, ensure_ascii=False, indent=2)
